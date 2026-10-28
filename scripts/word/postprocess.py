#!/usr/bin/env python3
"""Hau ky: muc luc tinh co so trang, so trang o le tren, trang bia rieng."""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PAGEMAP = Path("pagemap.json")
OUTNAME = "SO TAY DIEU TRA VIEN.docx"

# ---- Danh muc tieu de tu combined.md ----
entries = []
for line in Path("combined.md").read_text(encoding="utf-8").split("\n"):
    m = re.match(r"^(#{1,2})\s+(.+?)\s*$", line)
    if m:
        entries.append((len(m.group(1)), re.sub(r"[*`]", "", m.group(2)).strip()))

pages = []
if PAGEMAP.exists():
    pages = json.loads(PAGEMAP.read_text(encoding="utf-8"))
    assert len(pages) == len(entries), (len(pages), len(entries))

doc = Document("raw.docx")
body = doc.element.body

sdt = body.find(qn("w:sdt"))
assert sdt is not None, "khong tim thay truong muc luc"


def new_par(style=None, page_break=False):
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    if style:
        st = OxmlElement("w:pStyle"); st.set(qn("w:val"), style); ppr.append(st)
    if page_break:
        ppr.append(OxmlElement("w:pageBreakBefore"))
    p.append(ppr)
    return p


def add_text(p, text, bold=False, size=13):
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), "Times New Roman")
    rpr.append(rf)
    if bold:
        rpr.append(OxmlElement("w:b"))
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(size * 2)); rpr.append(sz)
    szc = OxmlElement("w:szCs"); szc.set(qn("w:val"), str(size * 2)); rpr.append(szc)
    r.append(rpr)
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text
    r.append(t)
    p.append(r)
    return r


def add_tab(p):
    r = OxmlElement("w:r")
    r.append(OxmlElement("w:tab"))
    p.append(r)


# ---- 1. Tieu de MUC LUC ----
head = new_par("TOCHeading", page_break=True)
add_text(head, "MỤC LỤC", bold=True, size=15)
sdt.addprevious(head)

# ---- 2. Cac dong muc luc tinh ----
for idx, (lvl, text) in enumerate(entries):
    p = new_par()
    ppr = p.find(qn("w:pPr"))
    # khong lui dau dong, gian dong gon
    ind = OxmlElement("w:ind")
    ind.set(qn("w:firstLine"), "0")
    ind.set(qn("w:left"), "0" if lvl == 1 else "340")
    ppr.append(ind)
    spc = OxmlElement("w:spacing")
    spc.set(qn("w:before"), "80" if lvl == 1 else "0")
    spc.set(qn("w:after"), "0")
    spc.set(qn("w:line"), "264"); spc.set(qn("w:lineRule"), "auto")
    ppr.append(spc)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right"); tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), "9070")
    tabs.append(tab); ppr.append(tabs)
    add_text(p, text, bold=(lvl == 1), size=13)
    add_tab(p)
    add_text(p, str(pages[idx]) if pages else "", bold=(lvl == 1), size=13)
    sdt.addprevious(p)

body.remove(sdt)

# ---- 3. So trang giua le tren, khong hien o trang bia ----
def page_field(par, size=13):
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.first_line_indent = Pt(0)
    par.paragraph_format.space_after = Pt(0)
    run = par.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), "Times New Roman")
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = " PAGE "
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    run._r.append(b); run._r.append(i); run._r.append(e)


for sec in doc.sections:
    sec.different_first_page_header_footer = True
    hdr = sec.header
    hdr.is_linked_to_previous = False
    page_field(hdr.paragraphs[0])
    fp = sec.first_page_header
    fp.is_linked_to_previous = False
    if fp.paragraphs:
        fp.paragraphs[0].text = ""

# ---- 4. Dinh dang bang: duong ke, lap tieu de, khong cat ngang dong ----
BORDER = ("top", "left", "bottom", "right", "insideH", "insideV")

def style_table(tb):
    tblpr = tb._tbl.tblPr
    for tag in ("w:tblBorders", "w:tblW", "w:tblLayout"):
        el = tblpr.find(qn(tag))
        if el is not None:
            tblpr.remove(el)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), "5000"); w.set(qn("w:type"), "pct")
    tblpr.append(w)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "autofit"); tblpr.append(lay)
    bd = OxmlElement("w:tblBorders")
    for side in BORDER:
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6")
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        bd.append(e)
    tblpr.append(bd)
    for ri, row in enumerate(tb.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for cell in row.cells:
            for par in cell.paragraphs:
                ppr = par._p.get_or_add_pPr()
                old_st = ppr.find(qn("w:pStyle"))
                if old_st is not None:
                    ppr.remove(old_st)
                st = OxmlElement("w:pStyle"); st.set(qn("w:val"), "TableText")
                ppr.insert(0, st)
                ind = ppr.find(qn("w:ind"))
                if ind is None:
                    ind = OxmlElement("w:ind"); ppr.append(ind)
                ind.set(qn("w:firstLine"), "0"); ind.set(qn("w:left"), "0")
                for r in par.runs:
                    r.font.size = Pt(11)
                if ri == 0:
                    for r in par.runs:
                        r.font.bold = True


for tb in doc.tables:
    style_table(tb)

doc.save(OUTNAME)
print(f"{OUTNAME}: {len(entries)} dòng mục lục"
      + (f", số trang {pages[0]}..{pages[-1]}" if pages else ", số trang để trống (lượt 1)"))
